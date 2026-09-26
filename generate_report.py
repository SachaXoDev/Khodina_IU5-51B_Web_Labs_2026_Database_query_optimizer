import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = docx.Document()

    # Поля страницы (2 см со всех сторон)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(0.79)
        section.right_margin = Inches(0.79)

    # Список 16 рисунков с подробными командами, URL, телами запросов и пояснениями
    figures = [
        # 1-10: Postman
        (
            "Рисунок 1 - GET списка услуг (с фильтром)",
            "Команда Postman:\n"
            "• Метод: GET\n"
            "• URL: http://localhost:3000/api/indexes?search=email&indexType=B-Tree\n"
            "• Параметры (Query Params): search = email, indexType = B-Tree\n"
            "• Ожидаемый ответ: 200 OK, JSON-массив опубликованных услуг.\n\n"
            "Пояснение: На бэкенде через TypeORM QueryBuilder выполняется фильтрация по подстроке поиска и типу индекса. Клиенту возвращаются исключительно опубликованные услуги; записи со статусом deleted и draft отфильтровываются в БД и клиенту не передаются (внутренний статус БД в JSON не возвращается)."
        ),
        (
            "Рисунок 2 - POST добавление новой услуги с картинкой и видео",
            "Команда Postman:\n"
            "• Метод: POST\n"
            "• URL: http://localhost:3000/api/indexes\n"
            "• Body (form-data):\n"
            "   - indexName (Text): idx_orders_status_brin\n"
            "   - shortDescription (Text): BRIN индекс для статусов заказов\n"
            "   - tableName (Text): orders\n"
            "   - indexType (Text): BRIN\n"
            "   - columnName (Text): status\n"
            "   - cardinality (Text): 500000\n"
            "   - fullDescription (Text): Оптимизация выборки по статусам заказов интернет-магазина\n"
            "   - image (File): файл изображения (.jpg / .png)\n"
            "   - video (File): файл короткого видео (.mp4)\n"
            "• Ожидаемый ответ: 201 Created, JSON созданного черновика услуги со всеми полями.\n\n"
            "Пояснение: Файлы изображения и короткого видео загружаются в MinIO. Названия файлов генерируются на латинице (например, index_image_...jpg и index_video_...mp4) и сохраняются в БД. Системные поля (внутренний статус 'draft' в БД, authorId = 1) вычисляются на бэкенде через функцию-singleton. Обязательным полем для черновика является только indexName, остальные поля опциональны, при этом клиенту возвращается консистентная структура всех полей (пустые поля заполнены значениями по умолчанию/null, без пропуска ключей)."
        ),
        (
            "Рисунок 3 - GET получение полей черновика",
            "Команда Postman:\n"
            "• Метод: GET\n"
            "• URL: http://localhost:3000/api/indexes/draft\n"
            "• Параметры: отсутствуют (ID в URL не передается)\n"
            "• Ожидаемый ответ: 200 OK, JSON черновика текущего пользователя.\n\n"
            "Пояснение: Идентификатор ID в запросе не передается по ТЗ. Бэкенд возвращает не более одной записи черновика для текущего зафиксированного пользователя (author_id = 1), вычисленного через функцию-singleton."
        ),
        (
            "Рисунок 4 - PUT публикация услуги",
            "Команда Postman:\n"
            "• Метод: PUT\n"
            "• URL: http://localhost:3000/api/indexes/:id/publish (где :id — идентификатор созданного черновика)\n"
            "• Headers: Content-Type: application/json\n"
            "• Body (raw JSON): {\"shortDescription\": \"Опубликованный оптимизированный индекс статусов\"}\n"
            "• Ожидаемый ответ: 200 OK, услуга опубликована, заполнено поле publishedAt.\n\n"
            "Пояснение: Выполняется смена статуса с 'draft' на 'published' в БД и фиксация системной даты publishedAt. Смена статуса возможна только создателем услуги и только из черновика; возврат опубликованной услуги обратно в статус черновика строго запрещен бизнес-логикой. В JSON-ответе клиенту внутренний статус не передается."
        ),
        (
            "Рисунок 5 - GET получение ленты без ид",
            "Команда Postman:\n"
            "• Метод: GET\n"
            "• URL: http://localhost:3000/api/indexes/feed\n"
            "• Параметры: отсутствуют\n"
            "• Ожидаемый ответ: 200 OK, JSON с полями current, prevId, nextId, hasPrev, hasNext.\n\n"
            "Пояснение: Возвращает первый опубликованный элемент ленты (current), а также вычисленные ссылки для циклической навигации: идентификатор предыдущего (prevId) и следующего (nextId) элементов ленты."
        ),
        (
            "Рисунок 6 - GET получение ленты по ид ?next=true",
            "Команда Postman:\n"
            "• Метод: GET\n"
            "• URL: http://localhost:3000/api/indexes/feed?id=1&next=true\n"
            "• Параметры (Query Params): id = 1, next = true\n"
            "• Ожидаемый ответ: 200 OK, следующий опубликованный элемент ленты.\n\n"
            "Пояснение: Осуществляет переход к следующей опубликованной услуге в ленте по отношению к переданному id."
        ),
        (
            "Рисунок 7 - POST постановка лайка",
            "Команда Postman:\n"
            "• Метод: POST\n"
            "• URL: http://localhost:3000/api/indexes/1/like\n"
            "• Headers: Content-Type: application/json\n"
            "• Body (raw JSON): {\"value\": 1}\n"
            "• Ожидаемый ответ: 201 Created, likesCount увеличен на 1, isLikedByCurrentUser = true.\n\n"
            "Пояснение: Лайк ставится от текущего пользователя (из функции-singleton). В таблице связей index_likes создается запись с отдельным первичным ключом, а счетчик likesCount у услуги атомарно увеличивается."
        ),
        (
            "Рисунок 8 - POST отмена лайка",
            "Команда Postman:\n"
            "• Метод: POST\n"
            "• URL: http://localhost:3000/api/indexes/1/like\n"
            "• Headers: Content-Type: application/json\n"
            "• Body (raw JSON): {\"value\": 0}\n"
            "• Ожидаемый ответ: 201 Created, likesCount уменьшен на 1, isLikedByCurrentUser = false.\n\n"
            "Пояснение: Значение 0 отменяет ранее поставленный лайк от текущего пользователя: запись из таблицы связей index_likes удаляется, счетчик likesCount уменьшается."
        ),
        (
            "Рисунок 9 - DELETE удаление услуги (soft delete)",
            "Команда Postman:\n"
            "• Метод: DELETE\n"
            "• URL: http://localhost:3000/api/indexes/:id\n"
            "• Ожидаемый ответ: 204 No Content (успешное мягкое удаление без тела ответа, стандарт REST).\n\n"
            "Пояснение: Реализовано мягкое удаление услуги: в БД статус меняется на 'deleted', физического удаления строки не происходит. Удаленная услуга перестает отдаваться клиенту в каталоге и ленте. В соответствии со стандартами REST, успешный DELETE-запрос возвращает HTTP 204 No Content без лишних полей в теле ответа."
        ),
        (
            "Рисунок 10 - POST регистрация нового пользователя",
            "Команда Postman:\n"
            "• Метод: POST\n"
            "• URL: http://localhost:3000/api/users/register\n"
            "• Headers: Content-Type: application/json\n"
            "• Body (raw JSON): {\"username\": \"ivan_architect\", \"password\": \"StrongPass123!\", \"role\": \"user\"}\n"
            "• Ожидаемый ответ: 201 Created, JSON зарегистрированного пользователя (поле password скрыто).\n\n"
            "Пояснение: Новый пользователь сохраняется в таблице users через ORM TypeORM с сохранением пароля в БД. Пароль скрыт от клиента в ответе с помощью сериализатора @Exclude()."
        ),

        # 11-13: SELECT в СУБД
        (
            "Рисунок 11 - SELECT из таблицы database_indexes",
            "SQL-запрос в Adminer / СУБД:\n"
            "SELECT id, index_name, status, image_url, video_url, author_id, likes_count, published_at FROM database_indexes ORDER BY id;\n\n"
            "Пояснение: Результат SQL-запроса в СУБД: отображаются сохраненные латинские имена файлов MinIO (image_url, video_url), статус опубликованной услуги (published), статус удаленной услуги (deleted), автор (author_id = 1) и дата публикации."
        ),
        (
            "Рисунок 12 - SELECT из таблицы index_likes",
            "SQL-запрос в Adminer / СУБД:\n"
            "SELECT id, user_id, index_id FROM index_likes ORDER BY id;\n\n"
            "Пояснение: Результат SQL-запроса в СУБД: таблица связей пользователь-услуга с отдельным первичным ключом id (user_id, index_id), подтверждающая фиксацию лайков от текущего пользователя."
        ),
        (
            "Рисунок 13 - SELECT из таблицы users",
            "SQL-запрос в Adminer / СУБД:\n"
            "SELECT id, username, password, role FROM users ORDER BY id;\n\n"
            "Пояснение: Результат SQL-запроса в СУБД: таблица пользователей системы, включая фиксированного создателя (pg_expert с id = 1), поле пароля password и нового зарегистрированного пользователя."
        ),

        # 14-16: Код
        (
            "Рисунок 14 - Функция-singleton с константой пользователя",
            "Файл проекта: database_query_optimizer/src/common/current-user.singleton.ts\n\n"
            "Пояснение: Класс CurrentUserSingleton с приватным конструктором и константой CURRENT_USER_ID = 1, а также функция-singleton getCurrentUserId(), возвращающая константу пользователя-создателя для всех методов сервиса."
        ),
        (
            "Рисунок 15 - Использование функции-singleton в методах сервиса",
            "Файл проекта: database_query_optimizer/src/indexes/indexes.service.ts\n\n"
            "Пояснение: Фрагменты кода методов getDraft(), createOrUpdateDraft(), publish(), toggleLike(), где текущий пользователь-создатель определяется через вызов функции-singleton getCurrentUserId()."
        ),
        (
            "Рисунок 16 - Модели Entity и DTO-сериализаторы с валидацией",
            "Файлы проекта: src/indexes/entities/database-index.entity.ts, src/indexes/dto/create-draft.dto.ts, src/indexes/dto/index-response.dto.ts\n\n"
            "Пояснение: Описание сущностей ORM, защита системных полей через запрет невалидированных свойств class-validator (forbidNonWhitelisted: true), скрытие паролей декоратором @Exclude() и компактная структура полей."
        ),
    ]

    for title, desc in figures:
        # Место под скриншот
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img = p_img.add_run("[Место для скриншота]")
        run_img.font.name = "Times New Roman"
        run_img.font.size = Pt(11)
        run_img.font.color.rgb = RGBColor(150, 150, 150)
        run_img.font.italic = True

        # Заголовок рисунка
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(4)
        p_title.paragraph_format.space_after = Pt(4)
        run_title = p_title.add_run(title)
        run_title.font.name = "Times New Roman"
        run_title.font.size = Pt(12)
        run_title.font.bold = True

        # Команда и пояснение
        p_desc = doc.add_paragraph()
        p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(16)

        # Выводим описание с сохранением переносов строк
        lines = desc.split("\n")
        for idx, line in enumerate(lines):
            run = p_desc.add_run(line)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            if line.startswith("Команда Postman:") or line.startswith("SQL-запрос") or line.startswith("Файл проекта:") or line.startswith("Пояснение:"):
                run.font.bold = True
            if idx < len(lines) - 1:
                p_desc.add_run("\n")

    out_path = "../WEB 3 лаба.docx"
    doc.save(out_path)
    print(f"Отчет успешно обновлен и сохранен в {out_path}")

if __name__ == "__main__":
    create_report()
